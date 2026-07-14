from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook

from app.agents.artifact_generation import ArtifactGenerationAgent, validate_workbook_terminology_and_stage_display
from app.io_utils import find_standard_json, load_process_definition
from app.tools.artifact_tools import generate_drawio_with_llm


class FakeArtifactLLM:
    def __init__(self, drawio_response: str | None = None) -> None:
        self.drawio_response = drawio_response

    def generate(self, system_prompt: str, user_prompt: str) -> str:
        if "xlsx skill 执行节点" in system_prompt:
            return """
import json
from pathlib import Path
from openpyxl import Workbook

def main(input_json_path, output_xlsx_path):
    data = json.loads(Path(input_json_path).read_text(encoding="utf-8"))
    wb = Workbook()
    ws = wb.active
    ws.title = "流程基本信息"
    ws.append([data["meta"]["process_name"] + " - 线上流程要素管理文档"])
    ws.append(["流程编号", data["meta"]["process_id"]])
    ws2 = wb.create_sheet("流程表单配置")
    ws2.append(["序号", "字段名称", "组件类型"])
    for field in data["form_fields"]:
        ws2.append([field["seq"], field["field_name"], field["component_type"]])
    ws3 = wb.create_sheet("流程环节配置")
    ws3.append(["序号", "环节名称", "处理角色"])
    for idx, node in enumerate(data["flow_nodes"], start=1):
        handler = node.get("handler") or {}
        ws3.append([idx, node["node_name"], handler.get("role", "")])
    wb.save(output_xlsx_path)

if __name__ == "__main__":
    import sys
    main(sys.argv[1], sys.argv[2])
"""
        if "draw.io" in system_prompt:
            return self.drawio_response or '<mxfile host="fake"><diagram name="流程图"><mxGraphModel><root><mxCell id="0"/></root></mxGraphModel></diagram></mxfile>'
        raise AssertionError(f"unexpected prompt: {system_prompt[:80]}")


def test_xlsx_tool_uses_skill_prompt_and_creates_workbook(tmp_path: Path) -> None:
    process = load_process_definition(find_standard_json("data/cases/EOA140_subsidiary_major_matter"))
    agent = ArtifactGenerationAgent(model=FakeArtifactLLM())

    result = agent.tools["generate_xlsx_from_process"](process, tmp_path / "process_readable.xlsx", tmp_path).payload

    assert Path(result["output_path"]).exists()
    assert Path(result["raw_output_path"]).exists()
    wb = load_workbook(result["output_path"])
    assert wb.sheetnames[:3] == ["流程基本信息", "流程表单配置", "流程环节配置"]
    assert wb["流程基本信息"]["A1"].value.endswith("线上流程要素管理文档")
    assert wb["流程表单配置"]["B2"].value == "标题"
    assert wb["流程表单配置"]["C2"].value == "单行文本"
    wb.close()
    report = json.loads((tmp_path / "xlsx_skill_report.json").read_text(encoding="utf-8"))
    assert report["status"] == "success"
    assert (tmp_path / "xlsx_agent_generated.py").exists()
    assert (tmp_path / "xlsx_agent_trace.json").exists()


def test_xlsx_skill_executes_generated_script_with_relative_run_dir(tmp_path: Path, monkeypatch) -> None:
    repo_root = Path(__file__).resolve().parents[1]
    monkeypatch.chdir(tmp_path)
    process = load_process_definition(find_standard_json(repo_root / "data/cases/EOA140_subsidiary_major_matter"))
    agent = ArtifactGenerationAgent(model=FakeArtifactLLM())

    result = agent.tools["generate_xlsx_from_process"](
        process,
        Path("relative_run/process_readable.xlsx"),
        Path("relative_run"),
    )

    assert result.status == "success"
    assert (tmp_path / "relative_run/process_readable.xlsx").exists()


def test_xlsx_terminology_validator_rejects_stage_node_and_raw_ids(tmp_path: Path) -> None:
    process = load_process_definition(find_standard_json("data/cases/EOA140_subsidiary_major_matter"))
    xlsx_path = tmp_path / "bad.xlsx"
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "流程表单配置"
    ws.append(["序号", "字段名称", "必填阶段", "可见阶段"])
    ws.append([1, "标题", "draft", "all"])
    ws2 = wb.create_sheet("流程环节配置")
    ws2.append(["节点ID", "节点名称", "目标节点ID"])
    ws2.append(["draft", "起草", "dept_supervisor"])
    wb.save(xlsx_path)

    report = validate_workbook_terminology_and_stage_display(xlsx_path, process)

    assert report["status"] == "terminology_or_stage_display_error"
    assert any(issue["type"] == "disallowed_term" for issue in report["issues"])
    assert any(issue["type"] == "raw_id_in_business_column" for issue in report["issues"])


def test_drawio_tool_accepts_fenced_xml(tmp_path: Path) -> None:
    process = load_process_definition(find_standard_json("data/cases/EOA140_subsidiary_major_matter"))
    response = """```xml
<mxfile host="fake"><diagram name="流程图"><mxGraphModel><root><mxCell id="0"/></root></mxGraphModel></diagram></mxfile>
```"""

    result = generate_drawio_with_llm(
        process,
        tmp_path / "process_flow.drawio",
        out_dir=tmp_path,
        model_or_provider=FakeArtifactLLM(response),
    )

    assert result["status"] == "success"
    assert (tmp_path / "process_flow.drawio").read_text(encoding="utf-8").startswith("<mxfile")
    report = json.loads((tmp_path / "drawio_validation_report.json").read_text(encoding="utf-8"))
    assert report["status"] == "success"


def test_drawio_tool_forces_black_font_color(tmp_path: Path) -> None:
    process = load_process_definition(find_standard_json("data/cases/EOA140_subsidiary_major_matter"))
    response = (
        '<mxfile host="fake"><diagram name="流程图"><mxGraphModel><root>'
        '<mxCell id="0"/>'
        '<mxCell id="1" value="&lt;font color=&quot;white&quot;&gt;审批&lt;/font&gt;" '
        'style="rounded=1;whiteSpace=wrap;html=1;fontColor=#FFFFFF;" vertex="1" parent="0">'
        '<mxGeometry width="120" height="60" as="geometry"/>'
        "</mxCell>"
        "</root></mxGraphModel></diagram></mxfile>"
    )

    result = generate_drawio_with_llm(
        process,
        tmp_path / "process_flow.drawio",
        out_dir=tmp_path,
        model_or_provider=FakeArtifactLLM(response),
    )

    output = (tmp_path / "process_flow.drawio").read_text(encoding="utf-8")
    assert result["status"] == "success"
    assert "fontColor=#000000" in output
    assert "#FFFFFF" not in output
    assert "white&quot;" not in output


def test_drawio_tool_reports_invalid_xml(tmp_path: Path) -> None:
    process = load_process_definition(find_standard_json("data/cases/EOA140_subsidiary_major_matter"))

    result = generate_drawio_with_llm(
        process,
        tmp_path / "process_flow.drawio",
        out_dir=tmp_path,
        model_or_provider=FakeArtifactLLM("<not-drawio />"),
    )

    assert result["status"] == "error"
    assert "mxfile" in result["message"]
    report = json.loads((tmp_path / "drawio_validation_report.json").read_text(encoding="utf-8"))
    assert report["status"] == "error"
