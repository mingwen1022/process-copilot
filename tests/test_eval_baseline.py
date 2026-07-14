from __future__ import annotations

import json
from pathlib import Path

import pytest

pytestmark = pytest.mark.skip(
    reason="eval 正在按 EOA140 单案例瘦身重构；本基线测试随新 scorer 一并重写"
)

from openpyxl import load_workbook

from app.eval.batch_eval import run_batch_eval
from app.eval.export_target_review import export_target_review
from app.eval.process_target_eval import evaluate_run_against_target, write_evaluation_report
from app.io_utils import load_process_definition, write_json
from app.reporting.comparison import compare_processes


def _write_run(run_dir: Path, process_payload: dict, clarifications: list[dict] | None = None) -> None:
    output_dir = run_dir / "workflow_design_output_writer"
    output_dir.mkdir(parents=True)
    write_json(
        {
            "status": "draft_ready",
            "process_definition": process_payload,
            "user_clarification_requests": clarifications or [],
        },
        output_dir / "workflow_design_output.json",
    )
    write_json(clarifications or [], output_dir / "user_clarification_requests.json")
    source_dir = run_dir / "source_ingestion"
    source_dir.mkdir()
    write_json([], source_dir / "source_file_manifest.json")


def test_target_review_exporter_outputs_expected_sheets(tmp_path: Path) -> None:
    paths = export_target_review(
        "data/01_leave_request/standard/leave_request_target.json",
        tmp_path / "review",
    )

    workbook = load_workbook(paths["xlsx"])
    assert {
        "Meta",
        "FormFields",
        "FlowNodes",
        "SubmitPaths",
        "Attachments",
        "CustomRoles",
        "Clarifications",
        "Guardrails",
        "EvidenceMap",
    }.issubset(set(workbook.sheetnames))
    assert paths["markdown"].read_text(encoding="utf-8").startswith("# Target Review:")


def test_comparison_normalizes_common_chinese_equivalents() -> None:
    expected = load_process_definition("data/01_leave_request/standard/leave_request_target.json")
    actual_data = expected.model_dump(mode="json")
    for node in actual_data["flow_nodes"]:
        for path in node["submit_paths"]:
            if path["condition"] and "请假天数>3天" in path["condition"]:
                path["condition"] = "结论性意见为同意，并且请假天数超过3天"
            if path["target_node_id"] == "DRAFT":
                path["path_name"] = "返回起草"
    actual = expected.__class__.model_validate(actual_data)

    report = compare_processes(actual, expected)

    assert report["path_condition_differences"] == []
    assert report["missing_submit_paths"] == []


def test_process_target_eval_outputs_review_status_and_scores(tmp_path: Path) -> None:
    target = Path("data/01_leave_request/standard/leave_request_target.json")
    process = load_process_definition(target)
    _write_run(
        tmp_path / "run",
        process.model_dump(mode="json"),
        clarifications=[
            {
                "id": "leave_proof_supplement_policy",
                "question": "病假证明是否允许返岗后补交？",
                "recommendation": "建议确认是否允许补交。",
                "options": ["允许补交", "不允许补交"],
            }
        ],
    )

    report = evaluate_run_against_target(tmp_path / "run", target)

    assert report["review_metadata"]["review_status"] == "draft"
    assert report["auto_score"]["total"] <= 100
    assert "human_adjusted_score" in report


def test_human_override_can_adjust_score(tmp_path: Path) -> None:
    target = Path("data/01_leave_request/standard/leave_request_target.json")
    process = load_process_definition(target)
    actual_data = process.model_dump(mode="json")
    actual_data["form_fields"] = [
        field for field in actual_data["form_fields"]
        if field["field_name"] != "代理人"
    ]
    run_dir = tmp_path / "run"
    _write_run(
        run_dir,
        actual_data,
        clarifications=[
            {"question": "病假证明是否允许返岗后补交？", "recommendation": "建议确认补交规则。"},
            {"question": "总经理审批请假天数门槛采用3天还是5天？", "recommendation": "建议确认最终门槛。"},
        ],
    )
    eval_dir = run_dir / "eval"
    eval_dir.mkdir()
    write_json(
        {
            "case_id": "01_leave_request",
            "reviewer": "ming",
            "overrides": [
                {
                    "item_type": "form_field",
                    "target_key": "代理人",
                    "actual_key": "",
                    "original_judgement": "missing",
                    "human_judgement": "match",
                    "reason": "测试 override 改分。",
                }
            ],
        },
        eval_dir / "human_review_overrides.json",
    )

    report = evaluate_run_against_target(run_dir, target)

    assert report["override_count"] == 1
    assert report["human_adjusted_score"]["total"] > report["auto_score"]["total"]


def test_batch_eval_writes_summary(tmp_path: Path) -> None:
    target = Path("data/01_leave_request/standard/leave_request_target.json")
    process = load_process_definition(target)
    run_dir = tmp_path / "run"
    _write_run(run_dir, process.model_dump(mode="json"))
    config = tmp_path / "eval_cases.json"
    write_json(
        {
            "cases": [
                {
                    "case_id": "01_leave_request",
                    "target": str(target),
                    "run": str(run_dir),
                }
            ]
        },
        config,
    )

    summary = run_batch_eval(config, tmp_path / "summary")

    assert summary["case_count"] == 1
    assert (tmp_path / "summary" / "eval_summary.json").exists()
    assert (tmp_path / "summary" / "eval_summary.csv").exists()
    assert (tmp_path / "summary" / "01_leave_request" / "evaluation_report.json").exists()


def test_write_evaluation_report_uses_expected_file_names(tmp_path: Path) -> None:
    report = {
        "case_id": "case",
        "status": "pass",
        "review_metadata": {"review_status": "draft"},
        "score": {"total": 100, "raw_total": 100, "deterministic_process_definition": 70, "clarification": 20, "evidence_and_guardrail": 10},
        "human_adjusted_score": {"total": 100},
        "override_count": 0,
        "hard_failures": [],
        "caps_applied": [],
        "comparison_report": {},
        "recommendations": [],
    }

    paths = write_evaluation_report(report, tmp_path)

    assert paths["json"].name == "evaluation_report.json"
    assert paths["markdown"].name == "evaluation_report.md"
