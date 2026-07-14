from __future__ import annotations

import csv
import json
import re
import zipfile
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from openpyxl import load_workbook


MAX_CHUNK_CHARS = 6_000
TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".json"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tiff", ".tif"}
BINARY_SIGNATURES = ("PK\x03\x04", "\x89PNG", "%PDF", "PNG IHDR")


@dataclass(frozen=True, slots=True)
class ParsedSource:
    manifest: dict[str, Any]
    chunks: list[dict[str, Any]]
    warnings: list[dict[str, Any]]


class _TextCollector(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self._skip_depth = 0
        self.parts: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag.lower() in {"script", "style"}:
            self._skip_depth += 1
        if tag.lower() in {"p", "div", "br", "tr", "li", "h1", "h2", "h3", "h4", "h5", "h6"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag.lower() in {"script", "style"} and self._skip_depth:
            self._skip_depth -= 1
        if tag.lower() in {"p", "div", "tr", "li"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self._skip_depth:
            return
        text = " ".join(data.split())
        if text:
            self.parts.append(text)

    def text(self) -> str:
        return _normalize_text(" ".join(self.parts))


def parse_source_file(
    path: str | Path,
    *,
    source_id: str,
    title: str | None = None,
    mime_type: str | None = None,
) -> ParsedSource:
    source_path = Path(path)
    source_type = detect_source_type(source_path, mime_type=mime_type)
    file_name = title or source_path.name
    warnings: list[dict[str, Any]] = []
    chunks: list[dict[str, Any]] = []

    try:
        if source_type in {"txt", "md"}:
            chunks = _chunks_for_text(_read_text(source_path), source_id=source_id, file_name=file_name, source_type=source_type)
        elif source_type == "json":
            chunks = _parse_json(source_path, source_id=source_id, file_name=file_name)
        elif source_type == "html":
            chunks = _parse_html(source_path, source_id=source_id, file_name=file_name, source_type="html")
        elif source_type == "docx":
            chunks = _parse_docx(source_path, source_id=source_id, file_name=file_name)
        elif source_type == "xlsx":
            chunks = _parse_xlsx(source_path, source_id=source_id, file_name=file_name)
        elif source_type == "csv":
            chunks = _parse_csv(source_path, source_id=source_id, file_name=file_name)
        elif source_type == "pdf":
            chunks, pdf_warnings = _parse_pdf_text(source_path, source_id=source_id, file_name=file_name)
            warnings.extend(pdf_warnings)
        elif source_type == "image":
            warnings.append(
                _warning(
                    source_id=source_id,
                    source_file=file_name,
                    warning_type="requires_vision_ocr",
                    message="图片 source 第一版暂不做 OCR，未进入 LLM 上下文。",
                )
            )
        else:
            warnings.append(
                _warning(
                    source_id=source_id,
                    source_file=file_name,
                    warning_type="unsupported_file_type",
                    message=f"暂不支持解析 {source_path.suffix or '未知'} 文件，未进入 LLM 上下文。",
                )
            )
    except Exception as exc:  # noqa: BLE001 - parser warnings should not abort the whole package.
        warnings.append(
            _warning(
                source_id=source_id,
                source_file=file_name,
                warning_type="parse_failed",
                message=f"解析失败：{exc}",
                level="blocking",
                repairable=True,
            )
        )
        chunks = []

    chunks = normalize_source_chunks(chunks)
    for chunk in chunks:
        if _contains_binary_signature(str(chunk.get("text") or "")):
            warnings.append(
                _warning(
                    source_id=source_id,
                    source_file=file_name,
                    warning_type="binary_text_filtered",
                    message="检测到二进制特征，已过滤该 chunk，避免进入 LLM 上下文。",
                    level="blocking",
                    repairable=True,
                )
            )
    chunks = [chunk for chunk in chunks if not _contains_binary_signature(str(chunk.get("text") or ""))]
    manifest = {
        "source_id": source_id,
        "file_name": file_name,
        "path": str(source_path),
        "source_type": source_type,
        "mime_type": mime_type,
        "size_bytes": source_path.stat().st_size if source_path.exists() else 0,
        "chunk_count": len(chunks),
        "chunk_ids": [chunk["chunk_id"] for chunk in chunks],
        "warnings": warnings,
    }
    return ParsedSource(manifest=manifest, chunks=chunks, warnings=warnings)


def detect_source_type(path: str | Path, *, mime_type: str | None = None) -> str:
    source_path = Path(path)
    suffix = source_path.suffix.lower()
    if suffix == ".docx":
        return "docx"
    if suffix in {".xlsx", ".xlsm"}:
        return "xlsx"
    if suffix == ".csv":
        return "csv"
    if suffix == ".pdf":
        return "pdf"
    if suffix in IMAGE_SUFFIXES:
        return "image"
    if suffix in TEXT_SUFFIXES:
        return suffix.removeprefix(".").replace("markdown", "md")
    if suffix == ".doc":
        sample = source_path.read_bytes()[:512].lstrip().lower()
        if sample.startswith(b"<!doctype html") or sample.startswith(b"<html"):
            return "html"
        return "doc"
    if mime_type:
        normalized = mime_type.lower()
        if "html" in normalized:
            return "html"
        if normalized.startswith("text/"):
            return "txt"
    sample = source_path.read_bytes()[:512].lstrip().lower()
    if sample.startswith(b"<!doctype html") or sample.startswith(b"<html"):
        return "html"
    return suffix.removeprefix(".") or "file"


def normalize_source_chunks(chunks: list[dict[str, Any]]) -> list[dict[str, Any]]:
    normalized: list[dict[str, Any]] = []
    for index, chunk in enumerate(chunks, start=1):
        text = _normalize_text(str(chunk.get("text") or ""))
        if not text:
            continue
        chunk_id = str(chunk.get("chunk_id") or f"{chunk.get('source_id', 'src')}_c{index:03d}")
        normalized.append(
            {
                "chunk_id": chunk_id,
                "source_id": chunk.get("source_id"),
                "source_file": chunk.get("source_file"),
                "source_type": chunk.get("source_type"),
                "location": chunk.get("location") or {"kind": "file"},
                "text": text,
                "char_count": len(text),
                "metadata": chunk.get("metadata") or {},
            }
        )
    return normalized


def _parse_json(path: Path, *, source_id: str, file_name: str) -> list[dict[str, Any]]:
    if path.name == "source_manifest.json":
        payload = json.loads(path.read_text(encoding="utf-8"))
        text = "Source manifest metadata:\n" + json.dumps(payload, ensure_ascii=False, indent=2)
        return _chunks_for_text(text, source_id=source_id, file_name=file_name, source_type="json")
    payload = json.loads(path.read_text(encoding="utf-8"))
    text = json.dumps(payload, ensure_ascii=False, indent=2)
    return _chunks_for_text(text, source_id=source_id, file_name=file_name, source_type="json")


def _parse_html(path: Path, *, source_id: str, file_name: str, source_type: str) -> list[dict[str, Any]]:
    collector = _TextCollector()
    collector.feed(_read_text(path))
    return _chunks_for_text(collector.text(), source_id=source_id, file_name=file_name, source_type=source_type)


def _parse_docx(path: Path, *, source_id: str, file_name: str) -> list[dict[str, Any]]:
    with zipfile.ZipFile(path) as archive:
        xml = archive.read("word/document.xml")
    root = ElementTree.fromstring(xml)
    paragraphs: list[str] = []
    namespace = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    for paragraph in root.iter(f"{namespace}p"):
        parts = [node.text or "" for node in paragraph.iter(f"{namespace}t")]
        text = _normalize_text("".join(parts))
        if text:
            paragraphs.append(text)
    return _chunks_for_text("\n".join(paragraphs), source_id=source_id, file_name=file_name, source_type="docx")


def _parse_xlsx(path: Path, *, source_id: str, file_name: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    chunks: list[dict[str, Any]] = []
    try:
        for sheet in workbook.worksheets:
            lines: list[str] = [f"Sheet: {sheet.title}"]
            row_count = 0
            for row in sheet.iter_rows(values_only=True):
                values = ["" if value is None else str(value).strip() for value in row]
                if not any(values):
                    continue
                row_count += 1
                lines.append("\t".join(values))
                if sum(len(line) for line in lines) >= MAX_CHUNK_CHARS:
                    chunks.extend(
                        _chunks_for_text(
                            "\n".join(lines),
                            source_id=source_id,
                            file_name=file_name,
                            source_type="xlsx",
                            location={"kind": "sheet", "sheet": sheet.title, "rows_read": row_count},
                            start_index=len(chunks) + 1,
                        )
                    )
                    lines = [f"Sheet: {sheet.title} (continued)"]
            if len(lines) > 1:
                chunks.extend(
                    _chunks_for_text(
                        "\n".join(lines),
                        source_id=source_id,
                        file_name=file_name,
                        source_type="xlsx",
                        location={"kind": "sheet", "sheet": sheet.title, "rows_read": row_count},
                        start_index=len(chunks) + 1,
                    )
                )
    finally:
        workbook.close()
    return chunks


def _parse_csv(path: Path, *, source_id: str, file_name: str) -> list[dict[str, Any]]:
    text = _read_text(path)
    rows = list(csv.reader(text.splitlines()))
    lines = ["\t".join(row) for row in rows if any(cell.strip() for cell in row)]
    return _chunks_for_text("\n".join(lines), source_id=source_id, file_name=file_name, source_type="csv")


def _parse_pdf_text(path: Path, *, source_id: str, file_name: str) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    data = path.read_bytes()
    decoded = data.decode("latin-1", errors="ignore")
    literals = re.findall(r"\(([^()]{4,300})\)", decoded)
    candidate = _normalize_text("\n".join(_decode_pdf_literal(item) for item in literals))
    chinese_count = len(re.findall(r"[\u4e00-\u9fff]", candidate))
    word_count = len(re.findall(r"[A-Za-z]{3,}", candidate))
    replacement_count = candidate.count("�")
    if chinese_count < 10 or replacement_count > 0 or chinese_count + word_count < 20:
        return [], [
            _warning(
                source_id=source_id,
                source_file=file_name,
                warning_type="requires_ocr",
                message="PDF 未抽取到足够正文，可能是扫描版；第一版暂不进入 LLM 上下文。",
            )
        ]
    return _chunks_for_text(candidate, source_id=source_id, file_name=file_name, source_type="pdf"), []


def _chunks_for_text(
    text: str,
    *,
    source_id: str,
    file_name: str,
    source_type: str,
    location: dict[str, Any] | None = None,
    start_index: int = 1,
) -> list[dict[str, Any]]:
    normalized = _normalize_text(text)
    if not normalized:
        return []
    chunks: list[dict[str, Any]] = []
    current: list[str] = []
    current_len = 0
    for paragraph in re.split(r"\n{2,}|\r\n", normalized):
        part = paragraph.strip()
        if not part:
            continue
        if current and current_len + len(part) + 2 > MAX_CHUNK_CHARS:
            chunk_index = start_index + len(chunks)
            chunks.append(
                _chunk(
                    source_id=source_id,
                    file_name=file_name,
                    source_type=source_type,
                    index=chunk_index,
                    text="\n".join(current),
                    location=location,
                )
            )
            current = []
            current_len = 0
        if len(part) > MAX_CHUNK_CHARS:
            for offset in range(0, len(part), MAX_CHUNK_CHARS):
                chunk_index = start_index + len(chunks)
                chunks.append(
                    _chunk(
                        source_id=source_id,
                        file_name=file_name,
                        source_type=source_type,
                        index=chunk_index,
                        text=part[offset : offset + MAX_CHUNK_CHARS],
                        location=location,
                    )
                )
            continue
        current.append(part)
        current_len += len(part) + 1
    if current:
        chunk_index = start_index + len(chunks)
        chunks.append(
            _chunk(
                source_id=source_id,
                file_name=file_name,
                source_type=source_type,
                index=chunk_index,
                text="\n".join(current),
                location=location,
            )
        )
    return chunks


def _chunk(
    *,
    source_id: str,
    file_name: str,
    source_type: str,
    index: int,
    text: str,
    location: dict[str, Any] | None,
) -> dict[str, Any]:
    return {
        "chunk_id": f"{source_id}_c{index:03d}",
        "source_id": source_id,
        "source_file": file_name,
        "source_type": source_type,
        "location": location or {"kind": "file", "file_name": file_name},
        "text": text,
        "metadata": {},
    }


def _warning(
    *,
    source_id: str,
    source_file: str,
    warning_type: str,
    message: str,
    level: str = "warning",
    repairable: bool = False,
) -> dict[str, Any]:
    return {
        "source_id": source_id,
        "source_file": source_file,
        "warning_type": warning_type,
        "level": level,
        "repairable": repairable,
        "message": message,
    }


def _read_text(path: Path) -> str:
    data = path.read_bytes()
    for encoding in ("utf-8", "utf-8-sig", "gb18030", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def _normalize_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[\x01-\x08\x0b\x0c\x0e-\x1f]", " ", text)
    lines = [" ".join(line.split()) for line in text.splitlines()]
    return "\n".join(line for line in lines if line).strip()


def _contains_binary_signature(text: str) -> bool:
    head = text[:500]
    return any(signature in head for signature in BINARY_SIGNATURES)


def _decode_pdf_literal(text: str) -> str:
    return text.replace("\\(", "(").replace("\\)", ")").replace("\\n", "\n").replace("\\r", "\n")
