from __future__ import annotations

import json
import subprocess
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable

from langchain_core.tools import StructuredTool
from openpyxl import load_workbook

from app.io_utils import write_json
from app.models.bedrock import create_bedrock_chat_model
from app.models.text import generate_text, is_legacy_text_generator, message_content_to_text
from app.tools.artifact_tools import generate_drawio_with_llm
from app.tracing import serialize_agent_response, serialize_llm_call, write_trace
from app.workflows.state import WorkflowState
from data.schema import ProcessDefinition


XLSX_SKILL_DIR = Path("/Users/ming/project/personal_assistant_agent/skills/xlsx")
XLSX_SKILL_PATH = XLSX_SKILL_DIR / "SKILL.md"
EXCEL_ERRORS = ("#VALUE!", "#DIV/0!", "#REF!", "#NAME?", "#NULL!", "#NUM!", "#N/A")


@dataclass(frozen=True, slots=True)
class ToolResult:
    name: str
    status: str
    payload: dict[str, Any]


class ArtifactGenerationAgent:
    """Agent node that owns the Excel and draw.io artifact tools."""

    def __init__(self, model: Any | None = None, agent: Any | None = None) -> None:
        self.model = model or create_bedrock_chat_model()
        self.tools: dict[str, Callable[[ProcessDefinition, Path, Path], ToolResult]] = {
            "generate_xlsx_from_process": self._generate_xlsx_from_process,
            "generate_drawio_from_process": self._generate_drawio_from_process,
        }
        self.agent = agent or (None if is_legacy_text_generator(self.model) else self._create_agent())

    def run(self, state: WorkflowState) -> WorkflowState:
        if not state.get("can_generate_artifacts") or state.get("candidate_process") is None:
            return {
                **state,
                "artifact_paths": {},
                "artifact_report": {"status": "skipped", "reason": "candidate_process failed schema validation"},
            }

        run_dir = Path(state["out_dir"]).resolve()
        agent_dir = run_dir / "artifact_generation_agent"
        agent_dir.mkdir(parents=True, exist_ok=True)
        process = state["candidate_process"]
        assert process is not None

        if self.agent is None:
            return self._run_owned_tools(state, process, run_dir, agent_dir)

        return self._run_langchain_agent(state, process, run_dir, agent_dir)

    def _run_owned_tools(
        self,
        state: WorkflowState,
        process: ProcessDefinition,
        run_dir: Path,
        agent_dir: Path,
    ) -> WorkflowState:
        """Run the same tools owned by this agent when the model is a plain text provider."""
        results: list[ToolResult] = []
        tool_plan = [
            (
                "generate_xlsx_from_process",
                agent_dir / "xlsx" / "process_readable.xlsx",
                agent_dir / "xlsx",
            ),
            (
                "generate_drawio_from_process",
                agent_dir / "drawio" / "process_flow.drawio",
                agent_dir / "drawio",
            ),
        ]
        for tool_name, output_path, tool_out_dir in tool_plan:
            try:
                results.append(self.tools[tool_name](process, output_path, tool_out_dir))
            except Exception as exc:
                results.append(
                    ToolResult(
                        name=tool_name,
                        status="error",
                        payload={"status": "error", "message": str(exc), "error_type": exc.__class__.__name__},
                    )
                )

        raw_output = "\n".join(f"{result.name}: {result.status}" for result in results)
        artifact_agent_raw = agent_dir / "artifact_agent_raw_output.txt"
        artifact_agent_raw.write_text(raw_output + "\n", encoding="utf-8")
        write_trace(
            {
                "agent": "artifact_generation_agent",
                "mode": "owned_tools",
                "tool_plan": [
                    {"tool_name": name, "output_path": str(output), "out_dir": str(tool_out_dir)}
                    for name, output, tool_out_dir in tool_plan
                ],
                "tool_results": {result.name: result.payload for result in results},
            },
            agent_dir / "messages.json",
        )
        if state.get("verbose"):
            _print_verbose_artifact_outputs(agent_dir, raw_output)

        artifact_paths = _expected_artifact_paths(run_dir)
        artifact_paths["artifact_agent_raw_output"] = str(artifact_agent_raw)
        artifact_report = _collect_artifact_report(run_dir, raw_output)
        artifact_report["tool_results"] = {result.name: result.payload for result in results}
        artifact_report["tool_statuses"] = {result.name: result.status for result in results}
        if any(result.status != "success" for result in results):
            artifact_report["status"] = "incomplete"
        return {**state, "artifact_paths": artifact_paths, "artifact_report": artifact_report}

    def _run_langchain_agent(
        self,
        state: WorkflowState,
        process: ProcessDefinition,
        run_dir: Path,
        agent_dir: Path,
    ) -> WorkflowState:
        user_prompt = _artifact_user_prompt(process.model_dump(mode="json"), agent_dir)
        try:
            response = self.agent.invoke(
                {
                    "messages": [
                        {
                            "role": "user",
                            "content": user_prompt,
                        }
                    ]
                }
            )
            raw_output = _extract_agent_text(response)
            write_trace(
                serialize_agent_response(
                    response,
                    input_messages=[
                        {"role": "system", "content": _artifact_system_prompt()},
                        {"role": "user", "content": user_prompt},
                    ],
                ),
                agent_dir / "messages.json",
            )
        except Exception as exc:
            raw_output = ""
            write_trace(
                {
                    "agent": "artifact_generation_agent",
                    "error_type": exc.__class__.__name__,
                    "message": str(exc),
                    "input_messages": [
                        {"role": "system", "content": _artifact_system_prompt()},
                        {"role": "user", "content": user_prompt},
                    ],
                },
                agent_dir / "messages.json",
            )
            artifact_paths = _expected_artifact_paths(run_dir)
            artifact_report = {
                "status": "error",
                "stage": "artifact_generation_agent",
                "message": str(exc),
                "files": {name: Path(path).exists() for name, path in artifact_paths.items()},
            }
            return {**state, "artifact_paths": artifact_paths, "artifact_report": artifact_report}

        artifact_agent_raw = agent_dir / "artifact_agent_raw_output.txt"
        artifact_agent_raw.write_text(raw_output, encoding="utf-8")
        if state.get("verbose"):
            print("[artifact_generation_agent] raw_output_begin")
            print(raw_output)
            print("[artifact_generation_agent] raw_output_end")

        artifact_paths = _expected_artifact_paths(run_dir)
        artifact_paths["artifact_agent_raw_output"] = str(artifact_agent_raw)
        artifact_report = _collect_artifact_report(run_dir, raw_output)
        return {**state, "artifact_paths": artifact_paths, "artifact_report": artifact_report}

    def _create_agent(self) -> Any:
        from langchain.agents import create_agent

        return create_agent(
            model=self.model,
            tools=self._build_langchain_tools(),
            system_prompt=_artifact_system_prompt(),
        )

    def _build_langchain_tools(self) -> list[StructuredTool]:
        def generate_xlsx_from_process(process_json: str, output_path: str, out_dir: str) -> str:
            process = ProcessDefinition.model_validate_json(process_json)
            result = self._generate_xlsx_from_process(process, Path(output_path), Path(out_dir))
            return json.dumps(result.payload, ensure_ascii=False)

        def generate_drawio_from_process(process_json: str, output_path: str, out_dir: str) -> str:
            process = ProcessDefinition.model_validate_json(process_json)
            result = self._generate_drawio_from_process(process, Path(output_path), Path(out_dir))
            return json.dumps(result.payload, ensure_ascii=False)

        return [
            StructuredTool.from_function(
                func=generate_xlsx_from_process,
                name="generate_xlsx_from_process",
                description=(
                    "Use the local xlsx skill instructions and local workspace command execution to generate "
                    "process_readable.xlsx. Arguments: process_json, output_path, out_dir."
                ),
            ),
            StructuredTool.from_function(
                func=generate_drawio_from_process,
                name="generate_drawio_from_process",
                description=(
                    "Generate an editable draw.io .drawio XML file from ProcessDefinition JSON. "
                    "Arguments: process_json, output_path, out_dir."
                ),
            ),
        ]

    def _generate_xlsx_from_process(self, process: ProcessDefinition, output_or_out_dir: Path, out_dir: Path) -> ToolResult:
        work_dir = _resolve_work_dir(out_dir)
        output_path = _resolve_output_path(output_or_out_dir, work_dir, "process_readable.xlsx")
        work_dir.mkdir(parents=True, exist_ok=True)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        raw_output_path = work_dir / "xlsx_skill_raw_output.txt"
        report_path = work_dir / "xlsx_skill_report.json"
        trace_path = work_dir / "xlsx_agent_trace.json"
        input_path = work_dir / "xlsx_skill_input.json"
        script_path = work_dir / "xlsx_agent_generated.py"
        process_payload = process.model_dump(mode="json")
        input_path.write_text(json.dumps(process_payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

        trace: dict[str, Any] = {
            "skill_path": str(XLSX_SKILL_PATH),
            "input_path": str(input_path),
            "script_path": str(script_path),
            "output_path": str(output_path),
            "iterations": [],
        }

        try:
            skill_text = _load_xlsx_skill()
            code = self._ask_llm_for_xlsx_code(skill_text, process_payload, output_path, work_dir)
            raw_outputs = [code]
            report = self._write_run_and_validate_xlsx(code, script_path, input_path, output_path, process, trace)

            if report["status"] != "success":
                repair_code = self._ask_llm_to_repair_xlsx_code(
                    skill_text=skill_text,
                    process_payload=process_payload,
                    previous_code=code,
                    validation_report=report,
                    trace=trace,
                    trace_dir=work_dir,
                )
                raw_outputs.append(repair_code)
                report = self._write_run_and_validate_xlsx(repair_code, script_path, input_path, output_path, process, trace)
        except Exception as exc:
            report = {
                "status": "error",
                "stage": "xlsx_skill_agent",
                "message": str(exc),
                "error_type": exc.__class__.__name__,
                "output_path": str(output_path),
            }
            raw_outputs = [report["message"]]

        raw_output_path.write_text("\n\n--- LLM OUTPUT ---\n\n".join(raw_outputs) + "\n", encoding="utf-8")
        report.update(
            {
                "stage": "xlsx_skill_agent",
                "xlsx_skill_dir": str(XLSX_SKILL_DIR),
                "xlsx_skill_path": str(XLSX_SKILL_PATH),
                "raw_output_path": str(raw_output_path),
                "trace_path": str(trace_path),
            }
        )
        write_json(report, report_path)
        write_json(trace, trace_path)
        payload = {
            "status": report["status"],
            "output_path": str(output_path),
            "raw_output_path": str(raw_output_path),
            "report_path": str(report_path),
            "trace_path": str(trace_path),
            "report": report,
        }
        return ToolResult("generate_xlsx_from_process", report["status"], payload)

    def _ask_llm_for_xlsx_code(
        self,
        skill_text: str,
        process_payload: dict[str, Any],
        output_path: Path,
        trace_dir: Path,
    ) -> str:
        system_prompt = _xlsx_codegen_system_prompt(skill_text)
        user_prompt = _xlsx_codegen_user_prompt(process_payload, output_path)
        raw_output = generate_text(self.model, system_prompt, user_prompt)
        write_trace(
            serialize_llm_call(
                agent="artifact_generation_agent",
                step="xlsx_codegen",
                system_prompt=system_prompt,
                user_prompt=user_prompt,
                assistant_output=raw_output,
                metadata={"output_path": str(output_path)},
            ),
            trace_dir / "codegen_messages.json",
        )
        return _extract_python_code(raw_output)

    def _ask_llm_to_repair_xlsx_code(
        self,
        *,
        skill_text: str,
        process_payload: dict[str, Any],
        previous_code: str,
        validation_report: dict[str, Any],
        trace: dict[str, Any],
        trace_dir: Path,
    ) -> str:
        system_prompt = _xlsx_codegen_system_prompt(skill_text)
        repair_prompt = f"""上一次生成的 Excel 代码没有通过校验。请基于 xlsx skill 和校验反馈，输出修正后的完整 Python 脚本。

要求仍然相同：只输出 Python 代码，不要 Markdown。

ProcessDefinition JSON:
{json.dumps(process_payload, ensure_ascii=False, indent=2)}

上一版代码:
```python
{previous_code}
```

校验报告:
{json.dumps(validation_report, ensure_ascii=False, indent=2)}

执行轨迹:
{json.dumps(trace.get("iterations", [])[-1:], ensure_ascii=False, indent=2)}
"""
        raw_output = generate_text(self.model, system_prompt, repair_prompt)
        write_trace(
            serialize_llm_call(
                agent="artifact_generation_agent",
                step="xlsx_repair",
                system_prompt=system_prompt,
                user_prompt=repair_prompt,
                assistant_output=raw_output,
                metadata={"previous_status": validation_report.get("status")},
            ),
            trace_dir / "repair_messages.json",
        )
        return _extract_python_code(raw_output)

    def _write_run_and_validate_xlsx(
        self,
        code: str,
        script_path: Path,
        input_path: Path,
        output_path: Path,
        process: ProcessDefinition,
        trace: dict[str, Any],
    ) -> dict[str, Any]:
        script_path.write_text(code.rstrip() + "\n", encoding="utf-8")
        if output_path.exists():
            output_path.unlink()

        command = [sys.executable, str(script_path), str(input_path), str(output_path)]
        completed = subprocess.run(
            command,
            cwd=str(script_path.parent),
            check=False,
            capture_output=True,
            text=True,
            timeout=90,
        )

        report: dict[str, Any] = {
            "status": "success" if completed.returncode == 0 and output_path.exists() else "error",
            "command": [Path(command[0]).name, str(script_path.name), str(input_path.name), str(output_path.name)],
            "returncode": completed.returncode,
            "stdout": completed.stdout,
            "stderr": completed.stderr,
            "output_path": str(output_path),
            "script_path": str(script_path),
        }

        if output_path.exists():
            workbook_report = scan_workbook_for_errors(output_path)
            content_report = validate_workbook_contains_process_content(output_path, process)
            terminology_report = validate_workbook_terminology_and_stage_display(output_path, process)
            inspection = inspect_xlsx_workbook(output_path)
            report.update(
                {
                    "workbook_scan": workbook_report,
                    "content_validation": content_report,
                    "terminology_validation": terminology_report,
                    "workbook_inspection": inspection,
                }
            )
            if report["status"] == "success" and workbook_report["status"] != "success":
                report["status"] = workbook_report["status"]
            if report["status"] == "success" and content_report["status"] != "success":
                report["status"] = content_report["status"]
            if report["status"] == "success" and terminology_report["status"] != "success":
                report["status"] = terminology_report["status"]
        else:
            report["message"] = "Excel output file was not created"

        trace["iterations"].append(report)
        return report

    def _generate_drawio_from_process(self, process: ProcessDefinition, output_or_out_dir: Path, out_dir: Path) -> ToolResult:
        work_dir = _resolve_work_dir(out_dir)
        output_path = _resolve_output_path(output_or_out_dir, work_dir, "process_flow.drawio")
        result = generate_drawio_with_llm(process, output_path, out_dir=work_dir, model_or_provider=self.model)
        return ToolResult("generate_drawio_from_process", result.get("status", "error"), result)


def _artifact_system_prompt() -> str:
    return """你是流程设计工作流的产物生成 Agent。

你拥有两个工具：
1. generate_xlsx_from_process：加载本地 xlsx skill 指令，写入并执行本次专用 Python，生成 process_readable.xlsx。
2. generate_drawio_from_process：生成 process_flow.drawio。

必须调用这两个工具。不要把 Excel 或 draw.io 做成 graph 节点；它们是当前 Agent 可调用的工具。"""


def _artifact_user_prompt(process_payload: dict[str, Any], out_dir: Path) -> str:
    process_json = json.dumps(process_payload, ensure_ascii=False, indent=2)
    xlsx_dir = out_dir / "xlsx"
    drawio_dir = out_dir / "drawio"
    return f"""请为以下 ProcessDefinition 生成两个产物。

输出目录：
{out_dir}

Excel 输出路径：
{xlsx_dir / "process_readable.xlsx"}
Excel tool out_dir：
{xlsx_dir}

draw.io 输出路径：
{drawio_dir / "process_flow.drawio"}
draw.io tool out_dir：
{drawio_dir}

ProcessDefinition JSON:
{process_json}
"""


def _xlsx_codegen_system_prompt(skill_text: str) -> str:
    return f"""你是 ArtifactGenerationAgent 内部的 xlsx skill 执行节点。

你已经加载本地 xlsx skill 指令，必须按这些指令生成 Excel 工作簿创建代码。

本地 xlsx skill 指令如下：
---
{skill_text}
---

你只能输出一个完整 Python 脚本，不要 Markdown、解释或代码块。
脚本必须满足：
1. 可用命令 `python <script> <input_json_path> <output_xlsx_path>` 执行。
2. 从第一个参数读取 ProcessDefinition JSON，从第二个参数写出 xlsx。
3. 只允许写入 output_xlsx_path，不要修改 raw data、standard data 或仓库代码。
4. 使用 openpyxl 或 pandas/openpyxl 生成 workbook；如无公式，不需要调用 LibreOffice recalc。
5. 不能生成空模板；必须从 JSON 动态写入全部表单字段、流程环节、角色、附件。
6. 生成文件至少包含这些 sheet：流程基本信息、流程表单配置、流程环节配置；若 JSON 有 roles/attachments，也要生成角色配置/附件说明。
7. 表单字段必须使用 schema 键名：seq、field_name、required_stages、visible_stages、editable_stages、component_type、logic_description、default_value、options、max_length。
8. 流程环节必须使用 schema 键名：node_id、node_name、is_draft、handler、opinion、opinion_label、time_limit_days、submit_paths。
9. handler 使用 mode/source/role/source_field；submit_paths 使用 path_name/condition/target_node_id。
10. Excel 可见文本必须统一业务叫法：用“环节”，不要用“阶段”或“节点”。例如表单字段列名用“必填环节 / 支持查看环节 / 支持编辑环节”，流程配置列名用“环节名称 / 处理人选择方式 / 提交路径 / 路径控制条件”。
11. Excel 可见文本必须中文为主。除“环节ID / 目标环节ID”等明确 ID 列外，不要直接显示 draft、all、dept_supervisor、financial_initial、DRAFT、END 等 raw id。
12. 环节引用展示为中文名称或“中文名称(id)”：draft 显示“起草(draft)”，all 显示“所有环节”，DRAFT 显示“退回起草(DRAFT)”，END 显示“流程结束(END)”，普通环节显示如“部门主管审批(dept_supervisor)”。
13. 字段叫法尽量对齐标准模板：流程编号、版本号、字段逻辑说明、选项/默认值、附件类型、上传环节、必传环节。
14. 文字、表头、列宽和冻结窗格要适合人工阅读，但具体排版由你根据 skill 和任务判断。
"""


def _xlsx_codegen_user_prompt(process_payload: dict[str, Any], output_path: Path) -> str:
    return f"""请根据下面的 ProcessDefinition JSON 生成本次专用的 Excel 创建脚本。

目标输出文件：
{output_path}

业务任务：
- 生成当前简化版流程 Excel，可读、可人工校对。
- 不要固化为固定空表；每个字段和流程环节都必须来自 JSON。
- 第一版重点保证信息完整：流程基本信息、表单字段、流程环节、路径条件、角色、附件。
- 如果字段列表、环节列表、角色或附件为空，就在对应位置写明“无”或跳过可选 sheet，不要留下看起来像数据缺失的空行。
- 表单字段配置里不得出现“阶段”字样；required_stages / visible_stages / editable_stages 展示列必须叫“必填环节 / 支持查看环节 / 支持编辑环节”。
- 流程环节配置里不得出现“节点”字样；node_id / target_node_id 若作为 ID 列展示，列名必须叫“环节ID / 目标环节ID”。更推荐同时展示“目标环节”，值用“环节中文名(id)”。
- 所有环节引用都中文为主：draft=起草(draft)，all=所有环节，DRAFT=退回起草(DRAFT)，END=流程结束(END)，其他 node_id 用 flow_nodes 里的 node_name 映射。

ProcessDefinition JSON:
{json.dumps(process_payload, ensure_ascii=False, indent=2)}
"""


def _load_xlsx_skill() -> str:
    if not XLSX_SKILL_PATH.is_file():
        raise FileNotFoundError(f"xlsx skill not found: {XLSX_SKILL_PATH}")
    return XLSX_SKILL_PATH.read_text(encoding="utf-8")


def _extract_python_code(raw_output: str) -> str:
    text = raw_output.strip()
    if text.startswith("```"):
        lines = text.splitlines()
        if lines and lines[0].startswith("```"):
            lines = lines[1:]
        if lines and lines[-1].strip() == "```":
            lines = lines[:-1]
        text = "\n".join(lines).strip()

    start_markers = ("from ", "import ", "#!")
    if not text.startswith(start_markers):
        for marker in start_markers:
            index = text.find(marker)
            if index != -1:
                text = text[index:].strip()
                break

    if "Workbook" not in text and "to_excel" not in text:
        raise ValueError("LLM output did not look like an Excel generation Python script")
    return text


def scan_workbook_for_errors(path: str | Path) -> dict[str, Any]:
    workbook_path = Path(path)
    wb = load_workbook(workbook_path, data_only=False)
    error_summary: dict[str, dict[str, Any]] = {}
    total_errors = 0
    total_formulas = 0

    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    value = cell.value
                    if isinstance(value, str) and value.startswith("="):
                        total_formulas += 1
                    if isinstance(value, str):
                        for error_token in EXCEL_ERRORS:
                            if error_token in value:
                                total_errors += 1
                                error_summary.setdefault(error_token, {"count": 0, "locations": []})
                                error_summary[error_token]["count"] += 1
                                error_summary[error_token]["locations"].append(f"{sheet_name}!{cell.coordinate}")
                                break
    finally:
        wb.close()

    return {
        "status": "success" if total_errors == 0 else "errors_found",
        "validator": "xlsx_workbook_scan",
        "total_errors": total_errors,
        "total_formulas": total_formulas,
        "error_summary": error_summary,
    }


def inspect_xlsx_workbook(path: str | Path, *, max_rows: int = 5, max_cols: int = 8) -> dict[str, Any]:
    workbook_path = Path(path)
    wb = load_workbook(workbook_path, data_only=False)
    try:
        sheets = []
        for ws in wb.worksheets:
            preview = []
            for row in ws.iter_rows(max_row=min(ws.max_row, max_rows), max_col=min(ws.max_column, max_cols), values_only=True):
                preview.append([value for value in row])
            sheets.append(
                {
                    "name": ws.title,
                    "max_row": ws.max_row,
                    "max_column": ws.max_column,
                    "preview": preview,
                }
            )
        return {"status": "success", "sheets": sheets}
    finally:
        wb.close()


def validate_workbook_contains_process_content(path: str | Path, process: ProcessDefinition) -> dict[str, Any]:
    workbook_path = Path(path)
    wb = load_workbook(workbook_path, data_only=False)
    cell_texts: list[str] = []
    try:
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        cell_texts.append(str(cell.value))
    finally:
        wb.close()

    joined = "\n".join(cell_texts)
    missing_form_fields = [field.field_name for field in process.form_fields if field.field_name not in joined]
    missing_flow_nodes = [node.node_name for node in process.flow_nodes if node.node_name not in joined]
    missing_meta = []
    if process.meta.process_name not in joined:
        missing_meta.append("meta.process_name")
    if process.meta.process_id not in joined:
        missing_meta.append("meta.process_id")

    status = "success" if not missing_form_fields and not missing_flow_nodes and not missing_meta else "content_missing"
    return {
        "status": status,
        "validator": "process_definition_content_presence",
        "missing_meta": missing_meta,
        "missing_form_fields": missing_form_fields,
        "missing_flow_nodes": missing_flow_nodes,
    }


def validate_workbook_terminology_and_stage_display(path: str | Path, process: ProcessDefinition) -> dict[str, Any]:
    workbook_path = Path(path)
    wb = load_workbook(workbook_path, data_only=False)
    raw_ids = {node.node_id for node in process.flow_nodes} | {"all", "DRAFT", "END"}
    issues: list[dict[str, Any]] = []

    try:
        for ws in wb.worksheets:
            headers = _infer_headers(ws)
            for row in ws.iter_rows():
                for cell in row:
                    value = cell.value
                    if not isinstance(value, str) or not value.strip():
                        continue
                    text = value.strip()
                    if "阶段" in text:
                        issues.append(
                            {
                                "sheet": ws.title,
                                "cell": cell.coordinate,
                                "type": "disallowed_term",
                                "message": "Excel 可见文本应使用“环节”，不要使用“阶段”",
                                "value": text,
                            }
                        )
                    if "节点" in text:
                        issues.append(
                            {
                                "sheet": ws.title,
                                "cell": cell.coordinate,
                                "type": "disallowed_term",
                                "message": "Excel 可见文本应使用“环节”，不要使用“节点”",
                                "value": text,
                            }
                        )

                    header = headers.get(cell.column, "")
                    if _looks_like_raw_stage_reference(text, raw_ids) and "ID" not in header.upper():
                        issues.append(
                            {
                                "sheet": ws.title,
                                "cell": cell.coordinate,
                                "type": "raw_id_in_business_column",
                                "message": "非 ID 列的环节引用必须中文为主，例如 起草(draft)、所有环节、流程结束(END)",
                                "header": header,
                                "value": text,
                            }
                        )
    finally:
        wb.close()

    return {
        "status": "success" if not issues else "terminology_or_stage_display_error",
        "validator": "workbook_terminology_and_stage_display",
        "issue_count": len(issues),
        "issues": issues,
    }


def _infer_headers(ws: Any) -> dict[int, str]:
    headers: dict[int, str] = {}
    for row in ws.iter_rows(max_row=min(ws.max_row, 5)):
        values = [cell.value for cell in row if isinstance(cell.value, str)]
        if len(values) < 2:
            continue
        if any(any(token in value for token in ("字段", "环节", "节点", "阶段", "路径", "附件")) for value in values):
            for cell in row:
                if isinstance(cell.value, str):
                    headers[cell.column] = cell.value.strip()
            break
    return headers


def _looks_like_raw_stage_reference(text: str, raw_ids: set[str]) -> bool:
    normalized = text.replace("，", ",").replace("；", ",").replace("、", ",").replace(" ", "")
    parts = [part for part in normalized.split(",") if part]
    if not parts:
        return False
    return all(part in raw_ids for part in parts)


def _resolve_work_dir(out_dir: Path) -> Path:
    return out_dir.resolve()


def _resolve_output_path(candidate: Path, work_dir: Path, default_name: str) -> Path:
    if candidate.suffix == ".xlsx" or candidate.suffix == ".drawio":
        return candidate.resolve() if candidate.is_absolute() else (Path.cwd() / candidate).resolve()
    return work_dir / default_name


def _extract_agent_text(response: Any) -> str:
    if isinstance(response, dict):
        messages = response.get("messages")
        if messages:
            return message_content_to_text(getattr(messages[-1], "content", messages[-1]))
        if isinstance(response.get("output"), str):
            return response["output"]
    return message_content_to_text(getattr(response, "content", response))


def _print_verbose_artifact_outputs(out_dir: Path, raw_output: str) -> None:
    print("[artifact_generation_agent] tool_status_begin")
    print(raw_output)
    print("[artifact_generation_agent] tool_status_end")
    for label, path in (
        ("xlsx_skill_raw_output", out_dir / "xlsx" / "xlsx_skill_raw_output.txt"),
        ("drawio_raw_output", out_dir / "drawio" / "drawio_raw_output.txt"),
    ):
        if path.exists():
            print(f"[artifact_generation_agent] {label}_begin")
            print(path.read_text(encoding="utf-8"))
            print(f"[artifact_generation_agent] {label}_end")


def _expected_artifact_paths(out_dir: Path) -> dict[str, str]:
    agent_dir = out_dir / "artifact_generation_agent"
    xlsx_dir = agent_dir / "xlsx"
    drawio_dir = agent_dir / "drawio"
    return {
        "artifact_agent_messages": str(agent_dir / "messages.json"),
        "artifact_agent_raw_output": str(agent_dir / "artifact_agent_raw_output.txt"),
        "process_readable": str(xlsx_dir / "process_readable.xlsx"),
        "xlsx_codegen_messages": str(xlsx_dir / "codegen_messages.json"),
        "xlsx_skill_input": str(xlsx_dir / "xlsx_skill_input.json"),
        "xlsx_skill_raw_output": str(xlsx_dir / "xlsx_skill_raw_output.txt"),
        "xlsx_agent_generated_script": str(xlsx_dir / "xlsx_agent_generated.py"),
        "xlsx_agent_trace": str(xlsx_dir / "xlsx_agent_trace.json"),
        "xlsx_skill_report": str(xlsx_dir / "xlsx_skill_report.json"),
        "process_flow": str(drawio_dir / "process_flow.drawio"),
        "drawio_messages": str(drawio_dir / "drawio_messages.json"),
        "drawio_raw_output": str(drawio_dir / "drawio_raw_output.txt"),
        "drawio_validation_report": str(drawio_dir / "drawio_validation_report.json"),
    }


def _collect_artifact_report(out_dir: Path, raw_output: str) -> dict[str, Any]:
    expected = _expected_artifact_paths(out_dir)
    existence = {name: Path(path).exists() for name, path in expected.items()}
    return {
        "status": "success" if all(existence.values()) else "incomplete",
        "files": existence,
        "agent_response": raw_output,
    }
