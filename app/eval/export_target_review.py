from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from app.io_utils import load_process_definition, load_standard_target


SHEETS = [
    "Meta",
    "FormFields",
    "FlowNodes",
    "SubmitPaths",
    "Attachments",
    "CustomRoles",
    "Clarifications",
    "Guardrails",
    "EvidenceMap",
]


def export_target_review(target_path: str | Path, out_dir: str | Path) -> dict[str, Path]:
    target = Path(target_path)
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)

    target_payload = load_standard_target(target)
    process = load_process_definition(target)
    workbook = Workbook()
    workbook.remove(workbook.active)

    _write_rows(
        workbook.create_sheet("Meta"),
        ["字段", "值"],
        [
            ["target_file", str(target)],
            ["review_status", _review_metadata(target_payload).get("review_status")],
            ["reviewer", _review_metadata(target_payload).get("reviewer")],
            ["reviewed_at", _review_metadata(target_payload).get("reviewed_at")],
            ["review_notes", _review_metadata(target_payload).get("review_notes")],
            ["process_id", process.meta.process_id],
            ["process_name", process.meta.process_name],
            ["version", process.meta.version],
            ["responsible_dept", process.meta.responsible_dept],
            ["description", process.meta.description],
            ["applicant_scope", process.meta.applicant_scope],
            ["entry_point", process.meta.entry_point],
        ],
    )
    _write_rows(
        workbook.create_sheet("FormFields"),
        ["序号", "字段名称", "组件类型", "必填环节", "可见环节", "可编辑环节", "选项/枚举值", "默认值", "逻辑说明", "最大长度"],
        [
            [
                field.seq,
                field.field_name,
                field.component_type.value,
                _join(field.required_stages),
                _join(field.visible_stages),
                _join(field.editable_stages),
                _join(field.options or []),
                field.default_value or "",
                field.logic_description or "",
                field.max_length or "",
            ]
            for field in process.form_fields
        ],
    )
    _write_rows(
        workbook.create_sheet("FlowNodes"),
        ["环节ID", "环节名称", "是否起草", "处理人选择方式", "处理人来源", "处理角色", "来源字段", "意见域名称", "结论性意见必填", "意见详情必填", "处理期限(天)"],
        [
            [
                node.node_id,
                node.node_name,
                "是" if node.is_draft else "否",
                node.handler.mode.value if node.handler else "",
                node.handler.source.value if node.handler else "",
                node.handler.role if node.handler else "",
                node.handler.source_field if node.handler else "",
                node.opinion_label or "",
                "" if node.opinion is None else ("是" if node.opinion.conclusive_required else "否"),
                "" if node.opinion is None else ("是" if node.opinion.detail_required else "否"),
                node.time_limit_days or "",
            ]
            for node in process.flow_nodes
        ],
    )
    _write_rows(
        workbook.create_sheet("SubmitPaths"),
        ["来源环节ID", "来源环节名称", "路径名称", "路径条件", "目标环节ID", "目标环节名称"],
        [
            [
                node.node_id,
                node.node_name,
                path.path_name,
                path.condition or "",
                path.target_node_id,
                process.get_node_name(path.target_node_id),
            ]
            for node in process.flow_nodes
            for path in node.submit_paths
        ],
    )
    _write_rows(
        workbook.create_sheet("Attachments"),
        ["附件类型", "上传环节", "必传环节"],
        [
            [item.attachment_type, _join(item.upload_stages), _join(item.required_stages)]
            for item in process.attachments or []
        ],
    )
    _write_rows(
        workbook.create_sheet("CustomRoles"),
        ["角色名称", "角色类型", "归属部门", "成员"],
        [
            [role.role_name, role.role_type.value, role.department or "", _join(role.members)]
            for role in process.roles or []
        ],
    )
    _write_rows(
        workbook.create_sheet("Clarifications"),
        ["ID", "问题", "期望关键词", "建议", "权重"],
        [
            [
                item.get("id", ""),
                item.get("question") or item.get("expected_question") or item.get("topic") or "",
                _join(item.get("expected_question_contains") or item.get("must_contain") or []),
                item.get("recommendation") or item.get("suggested_answer") or "",
                item.get("score_weight", ""),
            ]
            for item in target_payload.get("clarification_targets", []) or []
            if isinstance(item, dict)
        ],
    )
    _write_rows(
        workbook.create_sheet("Guardrails"),
        ["类型", "ID/名称", "说明"],
        _guardrail_rows(target_payload),
    )
    _write_rows(
        workbook.create_sheet("EvidenceMap"),
        ["评估项", "证据文件"],
        [
            [key, _join(value if isinstance(value, list) else [value])]
            for key, value in (target_payload.get("evidence_map") or {}).items()
        ],
    )

    xlsx_path = out / "target_review.xlsx"
    workbook.save(xlsx_path)
    md_path = out / "target_review.md"
    md_path.write_text(_render_markdown(target_payload, process), encoding="utf-8")
    return {"xlsx": xlsx_path, "markdown": md_path}


def _write_rows(sheet: Worksheet, headers: list[str], rows: list[list[Any]]) -> None:
    sheet.append(headers)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    for row in rows:
        sheet.append(row)
    for column in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 60)
        sheet.column_dimensions[column[0].column_letter].width = width
    sheet.freeze_panes = "A2"


def _render_markdown(target_payload: dict[str, Any], process: Any) -> str:
    lines = [
        f"# Target Review: {process.meta.process_name}",
        "",
        "## Meta",
        "",
        f"- review_status: {_review_metadata(target_payload).get('review_status')}",
        f"- process_id: {process.meta.process_id}",
        f"- process_name: {process.meta.process_name}",
        "",
        "## Counts",
        "",
        f"- form_fields: {len(process.form_fields)}",
        f"- flow_nodes: {len(process.flow_nodes)}",
        f"- submit_paths: {sum(len(node.submit_paths) for node in process.flow_nodes)}",
        f"- attachments: {len(process.attachments or [])}",
        f"- custom_roles: {len(process.roles or [])}",
        f"- clarifications: {len(target_payload.get('clarification_targets', []) or [])}",
        "",
        "## Form Fields",
        "",
    ]
    for field in process.form_fields:
        lines.append(f"- {field.seq}. {field.field_name} / {field.component_type.value} / options={_join(field.options or [])}")
    lines.extend(["", "## Flow Nodes", ""])
    for node in process.flow_nodes:
        handler = f"{node.handler.source.value}/{node.handler.role}" if node.handler else "-"
        lines.append(f"- {node.node_id}: {node.node_name} / handler={handler} / paths={len(node.submit_paths)}")
    lines.extend(["", "## Clarifications", ""])
    clarifications = target_payload.get("clarification_targets", []) or []
    if clarifications:
        for item in clarifications:
            lines.append(f"- {item.get('id')}: {item.get('question') or item.get('topic')}")
    else:
        lines.append("- 无")
    lines.append("")
    return "\n".join(lines)


def _guardrail_rows(target_payload: dict[str, Any]) -> list[list[Any]]:
    rows: list[list[Any]] = []
    for item in target_payload.get("forbidden_clarifications", []) or []:
        if isinstance(item, dict):
            rows.append(["forbidden_clarification", item.get("id", ""), item.get("description") or item.get("question") or ""])
    for item in target_payload.get("guardrail_targets", []) or []:
        if isinstance(item, dict):
            rows.append(["guardrail", item.get("id", ""), item.get("description") or item.get("rule") or ""])
        else:
            rows.append(["guardrail", "", str(item)])
    return rows


def _review_metadata(target_payload: dict[str, Any]) -> dict[str, Any]:
    metadata = target_payload.get("review_metadata") or {}
    return {
        "review_status": metadata.get("review_status") or "draft",
        "reviewer": metadata.get("reviewer"),
        "reviewed_at": metadata.get("reviewed_at"),
        "review_notes": metadata.get("review_notes") or "",
    }


def _join(values: list[Any]) -> str:
    return "、".join(str(item) for item in values if item is not None and str(item) != "")


def main() -> None:
    parser = argparse.ArgumentParser(description="Export target JSON to human review xlsx/md.")
    parser.add_argument("--target", required=True, help="Target JSON path.")
    parser.add_argument("--out", required=True, help="Output review directory.")
    args = parser.parse_args()

    paths = export_target_review(args.target, args.out)
    print(json.dumps({key: str(value) for key, value in paths.items()}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
